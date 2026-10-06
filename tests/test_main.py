"""
Focused unit and integration tests for main.py CLI entry point.

Verifies:
1. Single XML input with custom output path.
2. Default output filename (invoice_audit_report.xlsx) when omitted.
3. Folder containing XML files.
4. ZIP archive containing XML files.
5. Missing source path returns non-zero exit code (1).
6. Non-XML / unsupported source path returns non-zero exit code (1).
7. --help flag exits cleanly with status 0.
8. REVIEW invoices still produce a successful CLI run (exit code 0).
9. Corrupt/unparseable XML files are recorded in Excel without crashing the batch.
"""

from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import openpyxl

# Ensure project root is in sys.path
PROJECT_ROOT = (
    Path(__file__).resolve().parent.parent
    if Path(__file__).parent.name == "tests"
    else Path(__file__).resolve().parent
)
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from main import run_cli, entrypoint
from tests.test_batch import VALID_INVOICE_XML


class TestMainCLI(unittest.TestCase):

    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp(prefix="test_main_cli_"))

    def tearDown(self):
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_single_xml_custom_output(self):
        """CLI with single XML and custom output path succeeds with exit code 0."""
        xml_file = self.temp_dir / "single.xml"
        xml_file.write_text(VALID_INVOICE_XML.format(invoice_id="CLI-001"), encoding="utf-8")
        out_file = self.temp_dir / "my_custom_report.xlsx"

        exit_code = run_cli([str(xml_file), str(out_file)])
        self.assertEqual(exit_code, 0)
        self.assertTrue(out_file.exists())

        wb = openpyxl.load_workbook(out_file)
        self.assertEqual(wb["Özet"].max_row, 2)
        self.assertEqual(wb["Özet"].cell(row=2, column=2).value, "CLI-001")
        self.assertEqual(wb["Özet"].cell(row=2, column=1).value, "BAŞARILI")

    def test_default_output_filename(self):
        """CLI without output path creates 'invoice_audit_report.xlsx' in working directory."""
        xml_file = self.temp_dir / "invoice.xml"
        xml_file.write_text(VALID_INVOICE_XML.format(invoice_id="DEF-001"), encoding="utf-8")

        default_file = Path("invoice_audit_report.xlsx")
        try:
            exit_code = run_cli([str(xml_file)])
            self.assertEqual(exit_code, 0)
            self.assertTrue(default_file.exists())
            wb = openpyxl.load_workbook(default_file)
            self.assertIn("Özet", wb.sheetnames)
        finally:
            if default_file.exists():
                try:
                    default_file.unlink()
                except Exception:
                    pass

    def test_folder_input(self):
        """CLI with folder containing multiple XMLs succeeds."""
        (self.temp_dir / "inv1.xml").write_text(VALID_INVOICE_XML.format(invoice_id="DIR-1"), encoding="utf-8")
        (self.temp_dir / "inv2.xml").write_text(VALID_INVOICE_XML.format(invoice_id="DIR-2"), encoding="utf-8")
        out_file = self.temp_dir / "dir_report.xlsx"

        exit_code = run_cli([str(self.temp_dir), str(out_file)])
        self.assertEqual(exit_code, 0)
        self.assertTrue(out_file.exists())

        wb = openpyxl.load_workbook(out_file)
        self.assertEqual(wb["Özet"].max_row, 3)  # Header + 2 invoices

    def test_zip_input(self):
        """CLI with ZIP file containing XMLs succeeds."""
        zip_path = self.temp_dir / "bundle.zip"
        with zipfile.ZipFile(zip_path, "w") as zf:
            zf.writestr("z1.xml", VALID_INVOICE_XML.format(invoice_id="ZIP-1"))
            zf.writestr("z2.xml", VALID_INVOICE_XML.format(invoice_id="ZIP-2"))

        out_file = self.temp_dir / "zip_report.xlsx"
        exit_code = run_cli([str(zip_path), "-o", str(out_file)])
        self.assertEqual(exit_code, 0)
        self.assertTrue(out_file.exists())

        wb = openpyxl.load_workbook(out_file)
        self.assertEqual(wb["Özet"].max_row, 3)

    def test_missing_source_fails(self):
        """CLI with non-existent source returns exit code 1."""
        exit_code = run_cli(["non_existent_source_path_12345.xml"])
        self.assertEqual(exit_code, 1)

    def test_empty_or_non_xml_source_fails(self):
        """CLI with source having 0 XML files returns exit code 1."""
        empty_folder = self.temp_dir / "empty_dir"
        empty_folder.mkdir()
        exit_code = run_cli([str(empty_folder)])
        self.assertEqual(exit_code, 1)

        txt_file = self.temp_dir / "test.txt"
        txt_file.write_text("not an xml", encoding="utf-8")
        exit_code = run_cli([str(txt_file)])
        self.assertEqual(exit_code, 1)

    def test_help_flag_succeeds(self):
        """CLI with --help returns status 0."""
        exit_code = run_cli(["--help"])
        self.assertEqual(exit_code, 0)

    def test_review_invoice_exits_zero(self):
        """Invoices with REVIEW status do not trigger a CLI failure (exit code is 0)."""
        xml_dir = (
            PROJECT_ROOT / "test_data" / "xml"
            if (PROJECT_ROOT / "test_data" / "xml").exists()
            else PROJECT_ROOT
        )
        review_file = xml_dir / "test21.xml"
        out_file = self.temp_dir / "review_report.xlsx"

        exit_code = run_cli([str(review_file), str(out_file)])
        self.assertEqual(exit_code, 0)
        self.assertTrue(out_file.exists())

        wb = openpyxl.load_workbook(out_file)
        self.assertEqual(wb["Özet"].cell(row=2, column=1).value, "İNCELEME GEREKLİ")

    def test_error_invoice_handled_gracefully(self):
        """Batch containing corrupted XML records error in Excel and returns exit code 0."""
        valid_file = self.temp_dir / "good.xml"
        valid_file.write_text(VALID_INVOICE_XML.format(invoice_id="GOOD-01"), encoding="utf-8")

        corrupt_file = self.temp_dir / "broken.xml"
        corrupt_file.write_text("<Invoice><cbc:ID>broken unclosed xml", encoding="utf-8")

        out_file = self.temp_dir / "batch_with_error.xlsx"
        exit_code = run_cli([str(self.temp_dir), str(out_file)])
        self.assertEqual(exit_code, 0)
        self.assertTrue(out_file.exists())

        wb = openpyxl.load_workbook(out_file)
        ws_s = wb["Özet"]
        statuses = [ws_s.cell(row=r, column=1).value for r in range(2, ws_s.max_row + 1)]
        self.assertIn("BAŞARILI", statuses)
        self.assertIn("HATA", statuses)

    @patch("gui.launch_gui")
    def test_entrypoint_no_args_launches_gui(self, mock_launch):
        """Entrypoint with empty argument list dispatches to GUI."""
        exit_code = entrypoint([])
        self.assertEqual(exit_code, 0)
        mock_launch.assert_called_once()

    @patch("gui.launch_gui")
    def test_entrypoint_gui_flag_launches_gui(self, mock_launch):
        """Entrypoint with ['--gui'] flag dispatches to GUI."""
        exit_code = entrypoint(["--gui"])
        self.assertEqual(exit_code, 0)
        mock_launch.assert_called_once()

    @patch("gui.launch_gui")
    def test_entrypoint_help_does_not_launch_gui(self, mock_launch):
        """Entrypoint with ['--help'] shows help and does NOT launch GUI."""
        exit_code = entrypoint(["--help"])
        self.assertEqual(exit_code, 0)
        mock_launch.assert_not_called()


if __name__ == "__main__":
    unittest.main(verbosity=2)
